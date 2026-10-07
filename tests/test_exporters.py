import io

from docx import Document

from exporters import markdown_to_docx, transcriptions_to_docx


def _doc(buffer: io.BytesIO) -> Document:
    buffer.seek(0)
    return Document(buffer)


def test_markdown_to_docx_headings_and_lists():
    md = (
        "# Заголовок\n\n## Подзаголовок\n\n"
        "- пункт один\n- **жирный** пункт\n\n"
        "1. первый\n2. второй\n\nОбычный текст.\n"
    )
    doc = _doc(markdown_to_docx(md))

    texts = [p.text for p in doc.paragraphs]
    assert "Заголовок" in texts
    assert "Подзаголовок" in texts
    assert "пункт один" in texts
    assert "жирный пункт" in texts  # маркеры ** сняты
    assert "первый" in texts

    styles = {p.style.name for p in doc.paragraphs}
    assert "List Bullet" in styles
    assert "List Number" in styles


def test_markdown_to_docx_inline_formatting_stripped():
    md = "Текст с *курсивом* и _подчёркиванием_ и `кодом`."
    doc = _doc(markdown_to_docx(md))
    body = [p.text for p in doc.paragraphs if p.text.strip()]
    assert body == ["Текст с курсивом и подчёркиванием и `кодом`."]


def test_transcriptions_to_docx_structure():
    items = [
        {"filename": "call_1.wav", "transcript": "Привет, чем могу помочь?"},
        {"filename": "call_2.wav", "transcript": ""},
    ]
    doc = _doc(transcriptions_to_docx(items))
    full = "\n".join(p.text for p in doc.paragraphs)

    assert "call_1.wav" in full
    assert "Привет, чем могу помочь?" in full
    assert "call_2.wav" in full
    assert "Транскрибация не получена." in full  # пустой транскрипт

    # файл — это BytesIO с валидным docx (Document выше не упал)
    assert isinstance(transcriptions_to_docx(items), io.BytesIO)
