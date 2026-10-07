"""Экспорт результатов в документы Word (.docx)."""
from __future__ import annotations

import io
import os
import re
from datetime import datetime

from docx import Document
from docx.shared import Pt


def markdown_to_docx(markdown_text: str) -> io.BytesIO:
    """
    Конвертирует Markdown текст в Word документ (.docx).
    Возвращает BytesIO объект с содержимым документа.
    """
    doc = Document()

    # Устанавливаем стиль по умолчанию
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    lines = markdown_text.split('\n')
    i = 0

    while i < len(lines):
        line = lines[i].rstrip()

        # Пустая строка
        if not line:
            i += 1
            continue

        # Заголовок H1 (#)
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue

        # Заголовок H2 (##)
        if line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue

        # Заголовок H3 (###)
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        # Заголовок H4 (####)
        if line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)
            i += 1
            continue

        # Список (-, *, +)
        if line.lstrip().startswith(('- ', '* ', '+ ')):
            text = line.lstrip()[2:].strip()
            # Убираем markdown форматирование жирного и курсива
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            text = re.sub(r'__(.+?)__', r'\1', text)
            text = re.sub(r'_(.+?)_', r'\1', text)
            doc.add_paragraph(text, style='List Bullet')
            i += 1
            continue

        # Нумерованный список (1., 2., etc.)
        if re.match(r'^\s*\d+\.\s', line):
            text = re.sub(r'^\s*\d+\.\s+', '', line).strip()
            # Убираем markdown форматирование
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            doc.add_paragraph(text, style='List Number')
            i += 1
            continue

        # Обычный параграф
        text = line.strip()
        # Убираем markdown форматирование жирного и курсива
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        text = re.sub(r'\*(.+?)\*', r'\1', text)
        text = re.sub(r'__(.+?)__', r'\1', text)
        text = re.sub(r'_(.+?)_', r'\1', text)

        if text:
            doc.add_paragraph(text)

        i += 1

    # Сохраняем в BytesIO
    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)

    return docx_buffer


def transcriptions_to_docx(transcripts: list[dict[str, str]]) -> io.BytesIO:
    """
    Формирует Word-документ с блоками:
    название файла -> транскрибация -> пустая строка.
    """
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    doc.add_heading("Транскрибация диалогов", level=1)

    total = len(transcripts)
    for idx, item in enumerate(transcripts, start=1):
        filename = (item.get("filename") or "").strip() or "Без названия"
        transcript = (item.get("transcript") or "").strip()

        p_name = doc.add_paragraph()
        p_name.add_run("название файла: ").bold = True
        p_name.add_run(filename)

        p_transcript = doc.add_paragraph()
        p_transcript.add_run("транскрибация: ").bold = True
        p_transcript.add_run(transcript or "Транскрибация не получена.")

        # Отступ между файлами
        doc.add_paragraph("")
        # Явный разделитель между диалогами
        if idx < total:
            doc.add_paragraph("────────────────────────────────────────")
            doc.add_paragraph("")

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)
    return docx_buffer


_FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
     "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]


def _unicode_fonts() -> tuple[str, str] | None:
    for regular, bold in _FONT_CANDIDATES:
        if os.path.exists(regular) and os.path.exists(bold):
            return regular, bold
    return None


_MD_NOISE = re.compile(r"^#{1,6}\s+|^>\s+|^[-*+]\s+|[*_`~]+", flags=re.MULTILINE)


def chat_to_pdf(history: list[dict[str, str]]) -> bytes:
    """Рендерит историю чата в PDF (кириллица — через системный Unicode-шрифт)."""
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(16, 16, 16)

    fonts = _unicode_fonts()
    family = "helvetica"
    if fonts:
        pdf.add_font("body", style="", fname=fonts[0])
        pdf.add_font("body", style="B", fname=fonts[1])
        family = "body"

    pdf.add_page()
    pdf.set_font(family, "B", 16)
    pdf.set_text_color(80, 50, 180)
    pdf.cell(text="История чата", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(family, "", 9)
    pdf.set_text_color(140, 140, 140)
    pdf.cell(
        text=datetime.now().strftime("%d.%m.%Y %H:%M"),
        new_x="LMARGIN", new_y="NEXT",
    )
    pdf.ln(4)

    for msg in history:
        role = "Вы" if msg.get("role") == "user" else "Ассистент"
        text = _MD_NOISE.sub("", msg.get("content", "")).strip()
        text = re.sub(r"\n{3,}", "\n\n", text)
        if not text:
            continue

        pdf.set_font(family, "B", 10)
        pdf.set_text_color(80, 50, 180 if role == "Ассистент" else 60)
        pdf.cell(text=f"{role}:", new_x="LMARGIN", new_y="NEXT")

        pdf.set_font(family, "", 10)
        pdf.set_text_color(40, 40, 40)
        pdf.multi_cell(0, 5.5, text=text, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

    return bytes(pdf.output())
